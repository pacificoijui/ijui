#!/usr/bin/env python3
"""Converte a planilha REQUISIÇÕES em requisicao/dados/requisicoes.json.

    python3 requisicao/ferramentas/planilha-para-json.py PLANILHA.xlsx [--gravar]

Uma aba por secretaria, sempre as mesmas oito colunas (linha 1 = nome da
secretaria, linha 2 = títulos, dados da linha 3 em diante):

    REQUISIÇÃO | RECEBIDO EM | CREDOR | VALOR | OBJETO | MODALIDADE |
    EMPENHO | P/ CONTABILIDADE

Sem --gravar é ensaio: mostra o relatório e não toca em arquivo nenhum.
Também escreve requisicao/dados/CONFERIR.md com o que ficou em dúvida.
"""
import json, re, sys, unicodedata
from datetime import datetime, date
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parents[2]
DESTINO = RAIZ / "requisicao" / "dados" / "requisicoes.json"
avisos = []


def aviso(grau, secretaria, requisicao, campo, bruto, virou):
    avisos.append({"grau": grau, "sec": secretaria, "req": requisicao,
                   "campo": campo, "bruto": bruto, "virou": virou})


def texto(v):
    """Um traço sozinho é como a planilha escreve "não tem": vira vazio, para
    a tela não mostrar uma etiqueta de empenho com "-" dentro."""
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%d/%m/%Y")
    s = str(v).replace(" ", " ")
    s = unicodedata.normalize("NFC", " ".join(s.split()))
    return "" if s in ("-", "--", "---", "\u2013", "\u2014", ".") else s


def numero(v, sec, req):
    """Valor em reais. A planilha traz float na maioria das linhas e texto
    à moda brasileira em algumas ("25.000,00"); "-" quer dizer sem valor."""
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    s = texto(v)
    if s == "":                      # o traço da planilha já virou vazio
        return None
    limpo = re.sub(r"[^\d,.\-]", "", s)
    if limpo in ("", "-"):
        aviso("CONFERIR", sec, req, "valor", s, None)
        return None
    corte = max(limpo.rfind(","), limpo.rfind("."))
    casas = -1 if corte < 0 else len(limpo) - corte - 1
    try:
        if casas in (1, 2):
            n = float(limpo[:corte].replace(".", "").replace(",", "") + "." + limpo[corte + 1:])
        else:
            n = float(limpo.replace(".", "").replace(",", ""))
    except ValueError:
        aviso("CONFERIR", sec, req, "valor", s, None)
        return None
    aviso("ajustado", sec, req, "valor", s, n)
    return round(n, 2)


def data(v, sec, req, campo):
    """ISO aaaa-mm-dd. Vem como datetime na maioria das abas e como texto
    dd/mm/aaaa em SMCT e SMEL."""
    if v is None or v == "":
        return None
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    s = texto(v)
    if s == "":                      # idem: célula com um traço só
        return None
    m = re.match(r"^(\d{1,2})[/.-](\d{1,2})[/.-](\d{2,4})$", s)
    if m:
        d, mes, ano = (int(x) for x in m.groups())
        if ano < 100:
            ano += 2000
        try:
            return date(ano, mes, d).isoformat()
        except ValueError:
            pass
    aviso("CONFERIR", sec, req, campo, s, None)
    return None


def requisicao(v, sec):
    """O número da requisição, que a planilha guarda de três jeitos:

      "001/2026"        o normal;
      "002/2026 - B"    com sufixo, que é parte do número e não se perde;
      datetime(2026,7,1) o Excel leu "7/2026" como mês/ano e gravou o
                        primeiro dia daquele mês. O MÊS é o número da
                        requisição — confirmado pela vizinhança nas abas
                        (…006/2026, 2026-07-01, 008/2026…). Só acontece
                        até 12, que é onde o Excel ainda enxerga um mês.
    """
    if v is None or v == "":
        return None, None, "", ""
    if isinstance(v, (datetime, date)):
        num, ano = v.month, v.year
        rotulo = f"{num:03d}/{ano}"
        aviso("CONFERIR", sec, rotulo, "requisição",
              v.strftime("%d/%m/%Y") + " (virou data no Excel)", rotulo)
        return num, ano, "", rotulo
    s = texto(v)
    m = re.match(r"^(\d{1,4})\s*/\s*(\d{4})\s*(.*)$", s)
    if not m:
        aviso("CONFERIR", sec, s, "requisição", s, None)
        return None, None, "", s
    num, ano = int(m.group(1)), int(m.group(2))
    sufixo = m.group(3).strip(" -–—")
    rotulo = f"{num:03d}/{ano}" + (f" {sufixo}" if sufixo else "")
    return num, ano, sufixo, rotulo


def converter(caminho):
    wb = openpyxl.load_workbook(caminho, data_only=True)
    secretarias, itens, ident = [], [], 0
    for aba in wb.sheetnames:
        ws = wb[aba]
        sigla = texto(aba).strip()
        secretarias.append({"sigla": sigla, "nome": texto(ws["A1"].value) or sigla})
        for linha in ws.iter_rows(min_row=3, values_only=True):
            if not any(v not in (None, "", " ") for v in linha):
                continue
            num, ano, sufixo, rotulo = requisicao(linha[0], sigla)
            ident += 1
            itens.append({
                "id": ident,
                "sec": sigla,
                "num": num,
                "ano": ano,
                "sufixo": sufixo,
                "rotulo": rotulo,
                "recebido": data(linha[1], sigla, rotulo, "recebido em"),
                "credor": texto(linha[2]).upper(),
                "valor": numero(linha[3], sigla, rotulo),
                "objeto": texto(linha[4]),
                "modalidade": texto(linha[5]).upper(),
                "empenho": texto(linha[6]),
                "contabilidade": data(linha[7], sigla, rotulo, "p/ contabilidade"),
            })
    return secretarias, itens


def relatorio(secretarias, itens):
    print(f"\n{'='*66}\nPLANILHA DE REQUISIÇÕES -> JSON\n{'='*66}")
    print(f"secretarias ................ {len(secretarias)}")
    print(f"requisições ................ {len(itens)}")
    for s in secretarias:
        n = sum(1 for i in itens if i["sec"] == s["sigla"])
        print(f"  {s['sigla']:16} {n:5}   {s['nome'][:44]}")
    faltando = {c: sum(1 for i in itens if not i.get(c)) for c in
                ("num", "recebido", "credor", "valor", "objeto", "modalidade",
                 "empenho", "contabilidade")}
    print("\ncampos em branco:")
    for c, n in faltando.items():
        if n:
            print(f"  {c:16} {n}")
    conferir = [a for a in avisos if a["grau"] == "CONFERIR"]
    ajustes = [a for a in avisos if a["grau"] == "ajustado"]
    print(f"\n{'-'*66}\nPRECISA DE CONFERÊNCIA: {len(conferir)}\n{'-'*66}")
    for a in conferir[:40]:
        print(f"  {a['sec']:14} {a['req']:>14}  {a['campo']:16} {a['bruto']!r} -> {a['virou']!r}")
    if len(conferir) > 40:
        print(f"  ... e mais {len(conferir)-40} (todos no CONFERIR.md)")
    print(f"\nnormalizados sem dúvida: {len(ajustes)}")
    return conferir


def escrever_conferencia(conferir, itens):
    repetidos = {}
    for i in itens:
        if i["num"] is None:
            continue
        repetidos.setdefault((i["sec"], i["rotulo"]), []).append(i["id"])
    repetidos = {k: v for k, v in repetidos.items() if len(v) > 1}

    L = ["# Requisições a conferir\n",
         f"Gerado a partir da planilha em {datetime.now():%d/%m/%Y}. "
         f"São {len(conferir)} registros em que a planilha estava ambígua e o "
         f"conversor teve de decidir, mais {len(repetidos)} número(s) repetido(s) "
         f"dentro da mesma secretaria.\n"]
    if conferir:
        L += ["## Onde o conversor teve de decidir\n",
              "| Secretaria | Requisição | Campo | Estava na planilha | Entrou como |",
              "|---|---|---|---|---|"]
        for a in conferir:
            virou = a["virou"] if a["virou"] is not None else "em branco"
            L.append(f"| {a['sec']} | **{a['req']}** | {a['campo']} | `{a['bruto']}` | **{virou}** |")
    if repetidos:
        L += ["", f"## Número repetido na mesma secretaria ({len(repetidos)})\n",
              "A planilha traz mais de uma linha com o mesmo número. Entraram todas.\n",
              "| Secretaria | Requisição | Quantas |", "|---|---|---|"]
        for (sec, rot), ids in sorted(repetidos.items()):
            L.append(f"| {sec} | **{rot}** | {len(ids)} |")
    destino = RAIZ / "requisicao" / "dados" / "CONFERIR.md"
    destino.write_text("\n".join(L) + "\n", encoding="utf-8")
    print(f"relatório de conferência: {destino.relative_to(RAIZ)}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    secretarias, itens = converter(sys.argv[1])
    conferir = relatorio(secretarias, itens)
    if "--gravar" in sys.argv:
        escrever_conferencia(conferir, itens)
        saida = {"secretarias": secretarias,
                 "requisicoes": [{k: v for k, v in sorted(i.items())} for i in itens]}
        DESTINO.write_text(
            '{\n  "secretarias": [\n'
            + ",\n".join("    " + json.dumps(s, ensure_ascii=False) for s in saida["secretarias"])
            + '\n  ],\n  "requisicoes": [\n'
            + ",\n".join("    " + json.dumps(r, ensure_ascii=False) for r in saida["requisicoes"])
            + "\n  ]\n}\n", encoding="utf-8")
        print(f"\ngravado: {DESTINO.relative_to(RAIZ)}  ({len(itens)} requisições)")
    else:
        print("\n(ensaio — nada gravado. Use --gravar para valer.)")
